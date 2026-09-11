import json
import math

import pytest

from pcnn4.human_pref import (compute_weights, gen_comparisons, is_dupe,
                              load_groups, read_results)


@pytest.fixture(scope='module')
def comparisons():
    return gen_comparisons(seed=7)


def test_group_sizes(comparisons):
    groups = load_groups()
    assert sorted(groups) == [1, 2, 3, 4, 5, 6, 7]
    assert sum(len(v) for v in groups.values()) == 568


def test_comparison_counts(comparisons):
    groups = load_groups()
    by_pc = {}
    for c in comparisons:
        by_pc.setdefault(c['pc'], []).append(c)
    for pc, keys in groups.items():
        nd = sum(1 for k in keys if not is_dupe(k))
        dup = len(keys) - nd
        core = nd * (nd - 1) // 2
        if pc == 2:
            # PC2: dupes are valued manually, no dupe comparisons at all
            assert len(by_pc[pc]) == core == 595
        elif pc == 1:
            # PC1: each dupe gets 5 comparisons, topped up with other
            # dupes; deduped pairs can only reduce the total
            assert core == 0
            assert len(by_pc[pc]) <= dup * 5
            participation = {}
            for c in by_pc[pc]:
                for k in (c['a'], c['b']):
                    participation[k] = participation.get(k, 0) + 1
            for k in keys:
                if is_dupe(k):
                    assert participation.get(k, 0) >= 5
        else:
            assert len(by_pc[pc]) == core + dup * min(5, nd)


def test_comparisons_same_pc_and_deterministic(comparisons):
    for c in comparisons:
        assert len(c['a']) == len(c['b'])       # same segment size => same pc
        assert c['a'] != c['b']
    assert gen_comparisons(seed=7) == comparisons
    assert gen_comparisons(seed=8) != comparisons


def test_dupe_pairing_rule(comparisons):
    seen_pairs = set()
    for c in comparisons:
        pair = frozenset((c['pc'], c['a'], c['b']))
        assert pair not in seen_pairs      # no repeated matchup within a pc
        seen_pairs.add(pair)
        if c['pc'] == 1:
            continue                       # dupe-vs-dupe allowed (one nd only)
        if c['pc'] == 2:
            assert not is_dupe(c['a']) and not is_dupe(c['b'])
        else:
            assert not (is_dupe(c['a']) and is_dupe(c['b']))


def test_ids_unique(comparisons):
    ids = [c['id'] for c in comparisons]
    assert len(ids) == len(set(ids))


def test_weight_recovery_consistent_answers(comparisons):
    # answer PC3 (7 single-piece segments, complete graph) with perfectly
    # consistent ratios from known ground-truth weights -> recovered
    # weights must match up to scale
    truth = {'T': 100.0, 'I': 50.0, 'L': 10.0, 'J': 10.0,
             'S': 2.0, 'Z': 2.0, 'O': 1.0}
    answers = {}
    for c in comparisons:
        if c['pc'] == 3:
            answers[c['id']] = truth[c['a']] / truth[c['b']]
    results = compute_weights(comparisons, answers)
    res = results[3]
    assert res['answered'] == 21
    assert res['rms_log10'] < 1e-9
    w = res['weights']
    assert abs(w['T'] - 100.0) < 1e-6          # best scaled to 100
    for k in truth:
        assert abs(w[k] / w['O'] - truth[k]) < 1e-6
    # unanswered pcs list every segment with weight None (PC2 dupes land
    # here too, ready for manual values)
    assert results[1]['answered'] == 0
    assert set(results[1]['weights']) == set(load_groups()[1])
    assert all(w is None for w in results[1]['weights'].values())
    pc2 = results[2]['weights']
    assert len(pc2) == 140 and all(w is None for w in pc2.values())


def test_results_json_roundtrip(tmp_path):
    path = tmp_path / 'r.json'
    path.write_text(json.dumps({'answers': {'pc3-000': 5, 'pc3-001': 0.2}}))
    r = read_results(str(path))
    assert r == {'pc3-000': 5.0, 'pc3-001': 0.2}


def test_scaffold_roundtrip(tmp_path, comparisons):
    import openpyxl
    from pcnn4.human_pref import make_scaffold
    out = str(tmp_path / 's.xlsx')
    make_scaffold(comparisons[:5], out)
    wb = openpyxl.load_workbook(out)
    ws = wb['comparisons']
    ws.cell(row=2, column=5, value='5')
    ws.cell(row=3, column=5, value='1/20')
    wb.save(out)
    r = read_results(out)
    ids = [c['id'] for c in comparisons[:5]]
    assert r[ids[0]] == 5.0
    assert math.isclose(r[ids[1]], 0.05)
    assert len(r) == 2
